from __future__ import annotations

import csv
from pathlib import Path
from typing import List, Tuple

from ..i18n import t
from .ags4_import import looks_like_ags, read_ags
from .table_parser import ParsedData, parse_text
from .universal import read_grid, to_parsed

SUPPORTED_EXTENSIONS = (".xlsx", ".xlsm", ".xls", ".csv", ".txt", ".tsv",
                        ".ags")


def import_file(path: str) -> ParsedData:
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        sheets, error = _xlsx_sheets(p)
    elif suffix == ".xls":
        sheets, error = _xls_sheets(p)
    elif suffix in (".ods", ".numbers"):
        return _failure(p, t("import_warn_convert_xlsx"))
    else:
        text = _read_text(p)
        if suffix == ".ags" or looks_like_ags(text):
            return _combine(read_ags(text, source=p.name), p)
        sheets, error = [("", _csv_rows(text))], None
    if error:
        return _failure(p, error)

    found: List[ParsedData] = []
    for name, rows in sheets:
        source = f"{p.name} — {name}" if name and len(sheets) > 1 else p.name
        found.extend(to_parsed(read_grid(rows, name), source))
    if found:
        return _combine(found, p)

    lines = []
    for _, rows in sheets:
        for row in rows:
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                lines.append("  ".join(cells))
    return parse_text("\n".join(lines), source=p.name)


def _combine(samples: List[ParsedData], p: Path) -> ParsedData:
    samples = [s for s in samples if s.sieve]
    if not samples:
        return _failure(p, t("import_warn_too_few"))
    first = samples[0]
    first.others = samples[1:]
    return first


def _failure(p: Path, message: str) -> ParsedData:
    data = ParsedData(source=p.name)
    data.warnings.append(message)
    return data


def _read_text(p: Path) -> str:
    raw = p.read_bytes()
    for encoding in ("utf-8-sig", "cp1254", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _csv_rows(text: str):
    sample = text[:4096]
    try:
        sep = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        sep = ";" if sample.count(";") > sample.count(",") else ","
        if sample.count("\t") > max(sample.count(";"), sample.count(",")):
            sep = "\t"
    return [[cell.strip() for cell in row]
            for row in csv.reader(text.splitlines(), delimiter=sep)]


def _xlsx_sheets(p: Path) -> Tuple[list, str]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return [], t("import_warn_openpyxl")
    try:
        wb = load_workbook(p, read_only=True, data_only=True)
    except Exception:
        return [], t("import_warn_unreadable")
    sheets = []
    for ws in wb.worksheets:
        if getattr(ws, "sheet_state", "visible") != "visible":
            continue
        sheets.append((ws.title, [list(row) for row in ws.iter_rows(values_only=True)]))
    wb.close()
    return sheets, ""


def _xls_sheets(p: Path) -> Tuple[list, str]:
    try:
        import xlrd
    except ImportError:
        return [], t("import_warn_xlrd")
    try:
        book = xlrd.open_workbook(str(p))
    except Exception:
        return [], t("import_warn_unreadable")
    sheets = []
    for sheet in book.sheets():
        if getattr(sheet, "visibility", 0) != 0:
            continue
        rows = []
        for r in range(sheet.nrows):
            row = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_EMPTY or cell.ctype == xlrd.XL_CELL_BLANK:
                    row.append(None)
                elif cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        row.append(xlrd.xldate.xldate_as_datetime(cell.value,
                                                                  book.datemode))
                    except Exception:
                        row.append(cell.value)
                else:
                    row.append(cell.value)
            rows.append(row)
        sheets.append((sheet.name, rows))
    return sheets, ""
